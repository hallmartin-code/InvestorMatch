from __future__ import annotations

import io
import re

import pymupdf
import pytest
from openpyxl import load_workbook

from app.config import Settings
from app.errors import ExportUnavailableError
from app.models import CATEGORY_ORDER
from app.reports.excel_report import render_xlsx
from app.reports.frames import RANKED_COLUMNS, SHEET_NAMES, summary_counts
from app.reports.google_sheets import export_to_google_sheets
from app.reports.pdf_report import CONFIDENTIAL, check_layout, render_pdf, required_strings
from tests.conftest import investor, run


@pytest.fixture(scope="module")
def pdf(sample_result):
    return render_pdf(sample_result, "sample.xlsx")


@pytest.fixture(scope="module")
def workbook(sample_result):
    return load_workbook(io.BytesIO(render_xlsx(sample_result)))


def test_pdf_is_one_letter_page_without_clipping(pdf, sample_result):
    assert check_layout(pdf, required_strings(sample_result)) == []
    with pymupdf.open(stream=pdf, filetype="pdf") as doc:
        assert doc.page_count == 1
        page = doc[0]
        assert (round(page.rect.width), round(page.rect.height)) == (612, 792)
        sizes = {round(s["size"], 1) for b in page.get_text("dict")["blocks"] for l in b.get("lines", [])
                 for s in l["spans"] if s["text"].strip() and s["bbox"][1] < 740}
        assert min(sizes) >= 9.0      # body type ≥ 9 pt (only the footer band is smaller)
        assert CONFIDENTIAL in page.get_text()


def test_pdf_counts_match_workbook(pdf, workbook, sample_result):
    text = " ".join(pymupdf.open(stream=pdf, filetype="pdf")[0].get_text().split())
    ranked = list(workbook["Ranked Investors"].iter_rows(min_row=2, values_only=True))
    for cat in CATEGORY_ORDER:
        n = sum(1 for row in ranked if row[0] == cat.value)
        assert f"{cat.value} — {n} qualified" in text
    already = workbook["Already Introduced"].max_row - 1
    assert f"{already} investor(s) already introduced" in text
    review = sum(1 for row in workbook["Exclusions and Review"].iter_rows(min_row=2, values_only=True)
                 if row[0] == "REVIEW")
    assert summary_counts(sample_result)["review"] == review
    assert re.search(rf"Held for review {review}\b", text)


def test_pdf_never_pads_thin_categories(cfg):
    r = run(cfg, [investor()])
    pdf = render_pdf(r)
    text = pymupdf.open(stream=pdf, filetype="pdf")[0].get_text()
    assert text.count("No qualified contacts in this category.") == 3


def test_workbook_structure(workbook):
    assert workbook.sheetnames == SHEET_NAMES
    ws = workbook["Ranked Investors"]
    assert [c.value for c in ws[1]] == RANKED_COLUMNS
    for sheet in workbook:
        assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref
    assert all(ws.column_dimensions[col].width >= 9 for col in "ABCDEFGHIJKL")


def test_workbook_text_is_formula_safe(workbook):
    ws = workbook["Ranked Investors"]
    risky = [c for row in ws.iter_rows(min_row=2) for c in row if isinstance(c.value, str) and c.value[:1] in "=+-@"]
    assert risky, "sample contains formula-like values"
    for cell in risky:
        assert cell.data_type == "s" and cell.quotePrefix
    assert any(c.value == "=SUM(A1:A9) Capital" for c in risky)      # value itself unchanged


def test_missing_optional_fields_need_research(workbook):
    ws = workbook["Ranked Investors"]
    header = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        for name in ("Phone", "LinkedIn", "Location"):
            assert row[header.index(name)] not in (None, "")
        assert row[header.index("TEN Relationship (Y/N)")] in {"Y", "N", "needs research"}


class FakeWorksheet:
    def __init__(self, title):
        self.title, self.values, self.calls = title, None, []

    def update_title(self, title):
        self.title = title

    def resize(self, **kw):
        pass

    def update(self, values, start, value_input_option):
        self.values, self.option = values, value_input_option

    def freeze(self, rows):
        self.calls.append(("freeze", rows))

    def set_basic_filter(self, rng):
        self.calls.append(("filter", rng))

    def format(self, rng, fmt):
        pass


class FakeSpreadsheet:
    url = "https://docs.google.com/spreadsheets/d/fake"

    def __init__(self):
        self.sheet1 = FakeWorksheet("Sheet1")
        self.sheets = [self.sheet1]
        self.shared = []

    def add_worksheet(self, title, rows, cols):
        ws = FakeWorksheet(title)
        self.sheets.append(ws)
        return ws

    def share(self, address, perm_type, role, notify):
        self.shared.append((address, notify))


class FakeClient:
    def __init__(self):
        self.spreadsheet = FakeSpreadsheet()

    def create(self, title, folder_id=None):
        return self.spreadsheet


def test_google_sheets_export_with_fake_client(sample_result):
    client = FakeClient()
    settings = Settings(im_google_share_with="analyst@tencapital.example")
    url = export_to_google_sheets(sample_result, settings, "test", client=client)
    assert url.startswith("https://docs.google.com/")
    sheets = client.spreadsheet.sheets
    assert [s.title for s in sheets] == SHEET_NAMES
    assert all(s.option == "RAW" for s in sheets)
    assert sheets[0].values[0] == RANKED_COLUMNS
    assert len(sheets[0].values) - 1 == len(sample_result.ranked)
    assert client.spreadsheet.shared == [("analyst@tencapital.example", False)]   # never notifies by email


def test_google_sheets_quota_error_explains_shared_drive(sample_result):
    class QuotaClient:
        def create(self, title, folder_id=None):
            raise RuntimeError("APIError: [403]: The user's Drive storage quota has been exceeded.")

    with pytest.raises(ExportUnavailableError, match="Shared Drive"):
        export_to_google_sheets(sample_result, Settings(), "test", client=QuotaClient())


def test_google_sheets_unconfigured_explains_setup(sample_result):
    settings = Settings(im_google_service_account_file=None, google_application_credentials=None)
    with pytest.raises(ExportUnavailableError, match="IM_GOOGLE_SERVICE_ACCOUNT_FILE"):
        export_to_google_sheets(sample_result, settings, "test")
