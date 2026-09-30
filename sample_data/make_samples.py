"""Generate the SYNTHETIC sample inputs (no real companies, people or contact details).

    python sample_data/make_samples.py

Creates:
  cardiolyte_health_deck.pptx   12-slide deck: tables, speaker notes, footnotes, an image-only traction
                                slide (needs OCR), an appendix timeline with a conflicting close date,
                                and a future-round figure that must be ignored.
  cardiolyte_health_deck.pdf    the same deck as a PDF (image-only page included).
  TEN_Capital_Investor_List_SAMPLE.xlsx
                                ~90 contacts under non-standard headers with a title row above them, covering
                                duplicates, suppression, invalid/generic emails, excluded types, stage/geo
                                incompatibilities, ambiguous types, an organization with 5 contacts,
                                a formula-like value and a phone number starting with '+'.
  TEN_Intro_Tracker_2026.xlsx   multi-company intro tracker (with a typo'd company name)
  cardiolyte_health_intros_Q3.csv
                                company-specific intro list without a company column
All names use the reserved .example domain.
"""

from __future__ import annotations

import csv
import io
import random
from pathlib import Path

from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent
FONT_DIR = OUT.parent / "app" / "assets" / "fonts"

SLIDES: list[dict] = [
    {"title": "Cardiolyte Health", "body": ["Home cardiac rehab, prescribed like a drug", "Investor Presentation",
                                             "October 2026"]},
    {"title": "The Problem", "body": [
        "Only 24% of eligible heart patients complete cardiac rehabilitation¹",
        "Center-based programs require 36 in-person visits; most patients drop out by week 4",
        "¹ Footnote: synthetic statistic for demonstration purposes only."]},
    {"title": "Our Solution", "body": [
        "Cardiolyte is a prescription digital therapeutic (SaMD) for cardiac rehabilitation",
        "Remote patient monitoring with wearable ECG patch plus clinician dashboard",
        "Pursuing FDA De Novo clearance; cardiology-led clinical protocol"]},
    {"title": "Market", "body": [
        "TAM $6.1B: US cardiac rehab and cardiovascular digital health",
        "SAM $1.4B across 1,200 hospital cardiology programs"]},
    {"title": "Traction", "image_text": [
        "TRACTION", "3 health-system pilots in Texas", "412 patients enrolled", "81% program completion",
        "2 signed LOIs with cardiology groups"]},
    {"title": "Business Model", "body": [
        "$1,200 per patient program, reimbursed under cardiac rehab CPT codes",
        "B2B contracts with health systems and cardiology practices"]},
    {"title": "Team", "body": [
        "Headquartered in Austin, Texas",
        "CEO: Dana Whitfield (synthetic) — former VP Product, digital health scale-up",
        "CMO: Dr. Omar Reyes (synthetic) — interventional cardiologist"]},
    {"title": "Competition", "table": [
        ["", "Cardiolyte", "Center-based rehab", "Generic fitness apps"],
        ["Prescribed & reimbursed", "Yes", "Yes", "No"],
        ["Home-based", "Yes", "No", "Yes"],
        ["Clinician monitoring", "Yes", "Yes", "No"]]},
    {"title": "The Ask", "body": [
        "Raising $2.5M Seed round on a post-money SAFE ($12M valuation cap)",
        "Committed: $1.1M in signed SAFEs from angel investors",
        "Soft-circled: $400K (verbal interest from two family offices)",
        "Target close: December 15, 2026",
        "Seeking US-based investors who can co-invest alongside our lead angels",
        "Previously raised $750K pre-seed from friends & family (2025)"],
     "notes": "Speaker notes: walk through milestones; the round is open to co-investors."},
    {"title": "Use of Funds", "table": [
        ["Use of funds", "Share"], ["Clinical validation study", "40%"], ["Engineering", "30%"],
        ["Go-to-market", "20%"], ["Regulatory", "10%"]]},
    {"title": "Appendix: Milestone Timeline", "table": [
        ["Milestone", "Timing"], ["Seed close", "Q1 2027"], ["FDA De Novo submission", "Q3 2027"],
        ["Next round: Series A of $8M", "2028"]]},
    {"title": "Contact", "body": [
        "Cardiolyte Health, Inc. · 500 Example Ave, Austin, TX 78701",
        "founders@cardiolyte.example",
        "© 2026 Cardiolyte Health, Inc. Confidential."]},
]


def _font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    name = "OpenSans-Bold.ttf" if bold else "OpenSans-Regular.ttf"
    try:
        return ImageFont.truetype(str(FONT_DIR / name), size)
    except OSError:
        return ImageFont.load_default()


def traction_png(lines: list[str]) -> bytes:
    image = Image.new("RGB", (1600, 900), (244, 247, 251))
    draw = ImageDraw.Draw(image)
    y = 120
    for i, line in enumerate(lines):
        draw.text((120, y), line, fill=(20, 40, 80), font=_font(84 if i == 0 else 64, bold=i == 0))
        y += 150 if i == 0 else 120
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def make_pptx(path: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    for spec in SLIDES:
        slide = prs.slides.add_slide(blank)
        if "image_text" in spec:  # image-only slide: no text frames at all
            slide.shapes.add_picture(io.BytesIO(traction_png(spec["image_text"])), 0, 0,
                                     prs.slide_width, prs.slide_height)
            continue
        box = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(12), Inches(1))
        box.text_frame.text = spec["title"]
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(36)
        if "body" in spec:
            body = slide.shapes.add_textbox(Inches(0.7), Inches(1.7), Inches(12), Inches(5)).text_frame
            body.word_wrap = True
            for i, line in enumerate(spec["body"]):
                para = body.paragraphs[0] if i == 0 else body.add_paragraph()
                para.text = line
                para.runs[0].font.size = Pt(12 if line.startswith(("¹", "©")) else 20)
        if "table" in spec:
            rows, cols = len(spec["table"]), len(spec["table"][0])
            table = slide.shapes.add_table(rows, cols, Inches(0.7), Inches(1.7), Inches(12), Inches(0.5 * rows)).table
            for r, row in enumerate(spec["table"]):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
        if "notes" in spec:
            slide.notes_slide.notes_text_frame.text = spec["notes"]
    prs.save(path)


def make_pdf(path: Path) -> None:
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(TTFont("OpenSans", str(FONT_DIR / "OpenSans-Regular.ttf")))
    pdfmetrics.registerFont(TTFont("OpenSans-Bold", str(FONT_DIR / "OpenSans-Bold.ttf")))
    width, height = landscape(letter)
    c = canvas.Canvas(str(path), pagesize=(width, height))
    c.setTitle("Cardiolyte Health — Investor Presentation (synthetic)")
    for spec in SLIDES:
        if "image_text" in spec:
            c.drawImage(ImageReader(io.BytesIO(traction_png(spec["image_text"]))), 0, 0, width, height)
            c.showPage()
            continue
        c.setFont("OpenSans-Bold", 26)
        c.drawString(50, height - 70, spec["title"])
        y = height - 130
        for line in spec.get("body", []):
            c.setFont("OpenSans", 10 if line.startswith(("¹", "©")) else 15)
            c.drawString(50, y, line)
            y -= 30
        for row in spec.get("table", []):
            x = 50
            for value in row:
                c.setFont("OpenSans", 13)
                c.drawString(x, y, value)
                c.rect(x - 4, y - 8, 170, 26)
                x += 175
            y -= 26
        if "notes" in spec:
            pass  # PDFs carry no speaker notes
        c.showPage()
    c.save()


# ------------------------------------------------------------------------------ investor list

HEADERS = ["First", "Surname", "Firm", "E-mail Address", "Mobile", "Investor Category", "City/State",
           "Sector Interests", "Thesis Keywords", "Stage Preference", "Typical Check", "Geographic Focus",
           "Co-invest?", "LinkedIn Profile", "Email Status", "Core List", "Events Attended", "Notes"]

# (first, last, firm, email, phone, type, location, sectors, thesis, stage, check, geo, coinvest, linkedin,
#  status, core, events, notes)
CORE_ROWS: list[tuple] = [
    ("Avery", "Holt", "", "avery.holt@holtcapital.example", "+1 512 555 0142", "Angel Investor", "Austin, TX",
     "Digital Health, Cardiovascular", "digital therapeutics", "Pre-seed to Seed", "$25K–$100K", "Texas",
     "Yes", "https://www.linkedin.com/in/avery-holt-example", "Subscribed", "Yes", "TEN Summit 2025; Austin Pitch Night",
     "Former cardiologist"),
    ("Blake", "Moreno", "", "blake.moreno@morenofamily.example", "", "Angel", "Houston, TX", "Healthcare", "",
     "Seed", "$50K", "", "", "", "Subscribed", "", "2", ""),
    ("Casey", "Tran", "", "casey.tran@trancap.example", "", "Angel", "Dallas, TX", "Generalist", "", "Early stage",
     "", "", "", "", "", "", "", ""),
    ("Drew", "Kaplan", "", "drew.kaplan@kaplanangel.example", "", "Angel", "San Antonio, TX", "Fintech, Payments", "",
     "Seed", "$25K–$50K", "Texas", "", "", "", "", "", ""),
    ("Emery", "Sato", "Sato Holdings", "emery.sato@satoholdings.example", "", "", "Austin, TX",
     "Digital Health", "", "Seed", "$50K–$150K", "", "", "", "", "", "", "Type unknown"),
    ("Finley", "Ruiz", "Ruiz Partners", "finley.ruiz@ruizpartners.example", "", "Angel / VC", "Austin, TX",
     "Healthcare, Digital Therapeutics", "", "Seed", "$100K–$250K", "United States", "", "", "", "", "", ""),
    # duplicate pair: second copy (different case) is unsubscribed → both suppressed
    ("Gray", "Nolan", "", "gray.nolan@nolanangels.example", "", "Angel", "Austin, TX", "Digital Health", "",
     "Seed", "$25K–$75K", "Texas", "", "", "Subscribed", "", "", ""),
    ("Gray", "Nolan", "", "Gray.Nolan@NolanAngels.example", "", "Angel", "Austin, TX", "Digital Health", "",
     "Seed", "$25K–$75K", "Texas", "", "", "Unsubscribed", "", "", ""),
    # duplicate pair that merges: second copy adds phone + LinkedIn
    ("Harper", "Quinn", "", "harper.quinn@quinnhealth.example", "", "Angel", "Fort Worth, TX",
     "Cardiovascular, Medical Devices", "", "Seed", "$50K–$200K", "Southwest", "Yes", "", "Subscribed", "", "", ""),
    ("Harper", "Quinn", "", "HARPER.QUINN@quinnhealth.example", "+1 817 555 0199", "Angel", "Fort Worth, TX",
     "Cardiovascular", "", "Seed", "$50K–$200K", "Southwest", "Yes",
     "https://www.linkedin.com/in/harper-quinn-example", "Subscribed", "Core", "", ""),
    # invalid / missing emails
    ("Indigo", "Park", "", "indigo.park@@parkventures.example", "", "Angel", "Austin, TX", "Healthcare", "",
     "Seed", "", "", "", "", "", "", "", ""),
    ("Jules", "Bennett", "", "jules bennett at bennett.example", "", "Angel", "Austin, TX", "Healthcare", "",
     "Seed", "", "", "", "", "", "", "", ""),
    ("Kai", "Lindqvist", "", "", "", "Angel", "Austin, TX", "Healthcare", "", "Seed", "", "", "", "", "", "", "", ""),
    # angel groups
    ("", "", "Hill Country Angels", "info@hillcountryangels.example", "", "Angel Group", "Austin, TX",
     "Healthcare, Medical Devices, Digital Health", "", "Seed, Series A", "$100K–$500K", "Texas", "Yes", "",
     "Subscribed", "Yes", "TEN Summit 2025", "Generic inbox"),
    ("Logan", "Pierce", "Hill Country Angels", "logan.pierce@hillcountryangels.example", "", "Angel Group",
     "Austin, TX", "Healthcare, Digital Health", "", "Seed", "$100K–$500K", "Texas", "Yes", "", "Subscribed",
     "Yes", "", ""),
    ("Morgan", "Ellis", "Lone Star Health Syndicate", "morgan.ellis@lonestarhealth.example", "", "Syndicate",
     "Dallas, TX", "Digital Therapeutics, Cardiovascular", "", "Pre-seed, Seed", "$50K–$300K", "Texas, Southwest",
     "Yes", "https://www.linkedin.com/in/morgan-ellis-example", "", "", "", ""),
    ("Noel", "Vance", "Rust Belt Angel Network", "noel.vance@rustbeltangels.example", "", "Angel Network",
     "Cleveland, OH", "Healthcare, Digital Health", "", "Seed", "$50K–$250K", "Midwest", "", "", "", "", "", ""),
    ("Oakley", "Grant", "Bayou Health Angels", "oakley.grant@bayouhealth.example", "", "Angel Group",
     "New Orleans, LA", "Digital Health", "", "Seed", "$50K–$200K", "Gulf Coast", "", "", "", "", "", ""),
    # family offices / HNI
    ("Parker", "Sloan", "Meridian Family Office", "parker.sloan@meridianfo.example", "", "Family Office",
     "Houston, TX", "Healthcare, Cardiovascular", "", "Seed, Series A", "$250K–$1M", "United States", "Yes", "",
     "", "Yes", "", ""),
    ("Quincy", "Hale", "Crestview Single Family Office", "quincy.hale@crestviewsfo.example", "", "Single Family Office",
     "Denver, CO", "Generalist", "", "All stages", "$100K–$2M", "Global", "", "", "", "", "", ""),
    ("Reese", "Calder", "", "reese.calder@caldermail.example", "", "HNI", "Austin, TX", "Healthcare, Digital Health",
     "", "Seed", "$100K–$250K", "", "", "", "", "", "TEN Austin Dinner", ""),
    ("Sage", "Whitman", "Whitman Wealth", "sage.whitman@whitmanwealth.example", "", "Family Office", "Boston, MA",
     "Medical Devices", "", "Seed, Series A", "€100K–€500K", "North America", "", "", "", "", "", "EUR check sizes"),
    # VCs (Pulse Ventures has 5 contacts → cap at 3)
    ("Tatum", "Reid", "Pulse Ventures", "tatum.reid@pulseventures.example", "+1 512 555 0110", "Venture Capital",
     "Austin, TX", "Digital Health, Digital Therapeutics", "", "Seed", "$250K–$1.5M", "United States", "Yes",
     "https://www.linkedin.com/in/tatum-reid-example", "", "Yes", "TEN Summit 2025", ""),
    ("Uma", "Foster", "Pulse Ventures", "uma.foster@pulseventures.example", "", "Venture Capital", "Austin, TX",
     "Digital Health, Digital Therapeutics", "", "Seed", "$250K–$1.5M", "United States", "Yes", "", "", "Yes", "", ""),
    ("Vic", "Osei", "Pulse Ventures", "vic.osei@pulseventures.example", "", "Venture Capital", "Austin, TX",
     "Digital Health", "", "Seed", "$250K–$1.5M", "United States", "Yes", "", "", "", "", ""),
    ("Wren", "Adler", "Pulse Ventures", "wren.adler@pulseventures.example", "", "Venture Capital", "Austin, TX",
     "Digital Health", "", "Seed", "$250K–$1.5M", "United States", "", "", "", "", "", ""),
    ("Xander", "Bly", "Pulse Ventures", "xander.bly@pulseventures.example", "", "Venture Capital", "Austin, TX",
     "Healthcare", "", "Seed", "$250K–$1.5M", "", "", "", "", "", "", ""),
    ("Yael", "Corwin", "Medivance Corporate Ventures", "yael.corwin@medivance.example", "", "Corporate VC",
     "Minneapolis, MN", "Cardiovascular, Medical Devices", "", "Series A, Series B", "$2M–$5M", "United States",
     "", "", "", "", "", "Stage mismatch"),
    ("Zion", "Marsh", "Summit Growth Partners", "zion.marsh@summitgrowth.example", "", "Venture Capital",
     "New York, NY", "Healthcare", "", "Growth, Late stage", "$10M–$25M", "United States", "", "", "", "", "", ""),
    ("Adrian", "Stokes", "Granite Peak Private Equity", "adrian.stokes@granitepeak.example", "", "Private Equity",
     "Chicago, IL", "Healthcare services", "", "Buyout", "$25M+", "United States", "", "", "", "", "", ""),
    ("Bria", "Lang", "Harbor Point Securities", "bria.lang@harborpoint.example", "", "Investment Bank",
     "New York, NY", "Healthcare", "", "", "", "", "", "", "", "", "", ""),
    ("Cole", "Tanaka", "Keystone Seed Fund", "cole.tanaka@keystoneseed.example", "", "Seed Fund", "Pittsburgh, PA",
     "Digital Health, Healthcare IT", "", "Seed", "$2M–$4M", "United States", "", "", "", "", "", "Large min check"),
    ("Dara", "Whelan", "Solo Health Capital", "dara.whelan@solohealth.example", "", "Venture Capital",
     "Nashville, TN", "Digital Therapeutics", "", "Seed", "$500K–$1M", "United States", "Lead only", "", "", "",
     "", ""),
    ("Eli", "Novak", "Nordlicht Ventures", "eli.novak@nordlicht.example", "", "Venture Capital", "Berlin, Germany",
     "Digital Health", "", "Seed", "€250K–€1M", "Europe", "", "", "", "", "", ""),
    # prior-introduction targets
    ("Farah", "Iqbal", "Beacon Health Ventures", "farah.iqbal@beaconhealth.example", "", "Venture Capital",
     "Austin, TX", "Digital Health, Cardiovascular", "", "Seed", "$250K–$750K", "Texas", "Yes", "", "", "Yes", "",
     ""),
    ("Gideon", "Frost", "Frost Angel Partners", "gideon.frost@frostangels.example", "", "Angel", "Austin, TX",
     "Digital Therapeutics", "", "Seed", "$50K–$100K", "Texas", "", "", "", "", "", ""),
    ("Hana", "Okafor", "Okafor Family Office", "hana.okafor@okaforfo.example", "", "Family Office", "Dallas, TX",
     "Healthcare", "", "Seed", "$250K–$500K", "United States", "", "", "", "", "", ""),
    ("Ira", "Delgado", "Crossbridge Capital", "ira.delgado@crossbridge.example", "", "Venture Capital", "Houston, TX",
     "Digital Health", "", "Seed", "$250K–$1M", "United States", "", "", "", "", "", ""),
    ("Jonah", "Pryce", "Longhorn Angels", "jonah.pryce@longhornangels.example", "", "Angel Group", "Austin, TX",
     "Healthcare", "", "Seed", "$100K–$300K", "Texas", "", "", "", "", "", ""),
    # similar-client intro evidence (introduced to another healthcare company)
    ("Kira", "Voss", "", "kira.voss@vossangel.example", "", "Angel", "Austin, TX", "Digital Health", "",
     "Seed", "$25K–$100K", "", "", "", "", "", "", ""),
    # same person + org under two emails → flagged possible duplicate
    ("Lena", "Marsh", "Riverbend Family Office", "lena.marsh@riverbendfo.example", "", "Family Office",
     "Austin, TX", "Digital Health", "", "Seed", "$100K–$500K", "Texas", "", "", "", "", "", ""),
    ("Lena", "Marsh", "Riverbend Family Office", "lmarsh@riverbend-family.example", "", "Family Office",
     "Austin, TX", "Digital Health", "", "Seed", "$100K–$500K", "Texas", "", "", "", "", "", ""),
    # formula-like organization must stay text
    ("Milo", "Fenn", "=SUM(A1:A9) Capital", "milo.fenn@fenncap.example", "+1 737 555 0101", "Angel", "Austin, TX",
     "Digital Health", "", "Seed", "$25K–$50K", "Texas", "", "", "", "", "", "=HYPERLINK(\"http://x.example\")"),
]

FIRST = ["Aria", "Beck", "Cleo", "Dev", "Esme", "Flynn", "Gia", "Hugo", "Iris", "Jett", "Kaya", "Lior", "Mae",
         "Nico", "Ora", "Pax", "Rhea", "Soren", "Tess", "Ugo", "Vera", "Wade", "Xena", "Yuri", "Zara"]
LAST = ["Abbott", "Barrow", "Chen", "Dunmore", "Everly", "Fairbank", "Gould", "Hawthorne", "Ingram", "Jarvis",
        "Kessler", "Lowry", "Mercer", "Nash", "Oyelaran", "Prescott", "Quade", "Rourke", "Sutton", "Thorne"]
TYPES = ["Angel", "Angel Investor", "Angel Group", "Family Office", "HNWI", "Venture Capital", "Micro VC", "Seed Fund"]
CITIES = ["Austin, TX", "Houston, TX", "Dallas, TX", "San Antonio, TX", "Denver, CO", "Boston, MA", "Nashville, TN",
          "Phoenix, AZ", "Atlanta, GA", "Chicago, IL"]
SECTORS = ["Healthcare", "Digital Health", "Medical Devices, Diagnostics", "Digital Therapeutics", "Generalist",
           "Healthcare, Healthcare IT", "Cardiovascular, Digital Health", "Fintech", "Enterprise Software, SaaS",
           "Behavioral Health", ""]
STAGES = ["Seed", "Pre-seed, Seed", "Seed to Series A", "Early stage", "All stages", ""]
CHECKS = ["$25K–$50K", "$50K–$150K", "$100K–$500K", "$250K–$1M", "up to $100K", "", "$10K"]
GEOS = ["Texas", "United States", "Southwest", "", "North America", ""]


def investor_rows() -> list[list[str]]:
    rng = random.Random(20260930)
    rows = [list(r) for r in CORE_ROWS]
    for i in range(46):
        first, last = FIRST[i % len(FIRST)], LAST[(i * 7) % len(LAST)]
        kind = TYPES[rng.randrange(len(TYPES))]
        org = {"Angel": "", "Angel Investor": "", "HNWI": "", "Angel Group": f"{last} Angels",
               "Family Office": f"{last} Family Office"}.get(kind, f"{last} {rng.choice(['Ventures', 'Capital'])}")
        email = f"{first.lower()}.{last.lower()}{i}@{(org or last).lower().replace(' ', '')}.example"
        rows.append([first, last, org, email, "", kind, rng.choice(CITIES), rng.choice(SECTORS), "",
                     rng.choice(STAGES), rng.choice(CHECKS), rng.choice(GEOS), rng.choice(["Yes", "", ""]), "",
                     rng.choice(["Subscribed", "Subscribed", "", "Bounced"]) if i % 9 else "Unsubscribed",
                     rng.choice(["Yes", "", "", "No"]), rng.choice(["", "", "1", "TEN Summit 2025"]), ""])
    return rows


def make_investor_list(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Investors"
    ws.append(["TEN Capital Investor List — SYNTHETIC SAMPLE (not real people)"])
    ws.append(HEADERS)
    for row in investor_rows():
        ws.append(row)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"   # keep formula-like text as text in the source too
    wb.save(path)


def make_intro_lists(tracker: Path, company_csv: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Intros"
    ws.append(["Company", "Sector", "Investor", "Investor Email", "Firm", "Date Introduced", "Outcome"])
    ws.append(["Cardiolyte Health", "Digital Health", "Farah Iqbal", "FARAH.IQBAL@beaconhealth.example",
               "Beacon Health Ventures", "2026-08-04", "Passed — too early"])
    ws.append(["Cardiolyte Hlth", "Digital Health", "Hana Okafor", "hana.okafor@okaforfo.example",
               "Okafor Family Office", "2026-07-21", "Meeting held"])
    ws.append(["VitalSync Medical", "Healthcare", "Kira Voss", "kira.voss@vossangel.example", "", "2026-03-02",
               "Invested"])
    ws.append(["LedgerLoop", "Fintech", "Casey Tran", "casey.tran@trancap.example", "", "2026-02-11", "Passed"])
    ws.append(["Cardiolyte Health", "Digital Health", "Rowan Ashby", "rowan.ashby@ashbycap.example",
               "Ashby Capital", "2026-08-19", "No response"])
    wb.save(tracker)

    with company_csv.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["Investor Name", "Organization", "Intro Date", "Status"])
        writer.writerow(["Gideon Frost", "Frost Angel Partners", "2026-09-02", "Intro sent"])
        writer.writerow(["Ira Delgado", "Delgado Family Trust", "2026-09-05", "Intro sent"])
        writer.writerow(["", "Bayou Health Angels", "2026-09-09", "Group intro via portal"])
        writer.writerow(["Jonah Pryce", "Longhorn Angels", "2026-09-12", "Meeting scheduled"])


if __name__ == "__main__":
    make_pptx(OUT / "cardiolyte_health_deck.pptx")
    make_pdf(OUT / "cardiolyte_health_deck.pdf")
    make_investor_list(OUT / "TEN_Capital_Investor_List_SAMPLE.xlsx")
    make_intro_lists(OUT / "TEN_Intro_Tracker_2026.xlsx", OUT / "cardiolyte_health_intros_Q3.csv")
    print("Synthetic samples written to", OUT)
