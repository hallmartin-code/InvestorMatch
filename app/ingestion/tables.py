"""CSV / Excel readers that keep every cell as text and remember file, sheet and row numbers."""

from __future__ import annotations

import csv
import re
import hashlib
import io
from dataclasses import dataclass, field
from pathlib import Path

from app.errors import IngestError, UnsupportedFormatError
from app.utils.text import clean_cell

TABLE_EXT = {".csv", ".tsv", ".txt", ".xlsx", ".xlsm", ".xls"}
HEADER_SCAN_ROWS = 15
_EMAIL_CELL = re.compile(r"^[^@\s]+@[^@\s]+\.[A-Za-z]{2,}$")


@dataclass
class ImportedTable:
    filename: str
    sheet: str | None
    columns: list[str]
    rows: list[dict[str, str]]
    row_numbers: list[int]             # 1-based row number in the source file/sheet
    sha256: str = ""
    warnings: list[str] = field(default_factory=list)

    @property
    def label(self) -> str:
        return f"{self.filename} › {self.sheet}" if self.sheet else self.filename

    def sample(self, column: str, limit: int = 3) -> list[str]:
        return [r[column] for r in self.rows if r.get(column)][:limit]


def _header_index(grid: list[list[str]]) -> int:
    """First row (within the first rows) with at least two non-empty cells: skips title rows."""
    for index, row in enumerate(grid[:HEADER_SCAN_ROWS]):
        if sum(1 for c in row if c) >= 2:
            return index
    return 0


def _unique(columns: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for index, name in enumerate(columns):
        name = name or f"Column {index + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        out.append(name)
    return out


def _from_grid(grid: list[list[str]], filename: str, sheet: str | None, sha: str) -> ImportedTable | None:
    grid = [row for row in grid]
    if not any(any(c for c in row) for row in grid):
        return None
    header = _header_index(grid)
    header_cells = [clean_cell(c) for c in grid[header]]
    headerless = any(_EMAIL_CELL.match(c) for c in header_cells if c)
    if headerless:  # the first row is data (an email address is never a column header)
        width = max(len(r) for r in grid)
        columns = [f"Column {i + 1}" for i in range(width)]
        data_start = header
    else:
        columns = _unique(header_cells)
        data_start = header + 1
    rows, numbers = [], []
    for offset, raw in enumerate(grid[data_start:], start=data_start + 1):
        cells = [clean_cell(c) for c in raw] + [""] * max(0, len(columns) - len(raw))
        if not any(cells):
            continue
        rows.append(dict(zip(columns, cells[: len(columns)], strict=False)))
        numbers.append(offset)
    table = ImportedTable(filename=filename, sheet=sheet, columns=columns, rows=rows, row_numbers=numbers, sha256=sha)
    if headerless:
        table.warnings.append("No header row found (the first row holds data); columns are named Column 1, 2, …")
    return table


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
            if encoding == "utf-16" and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
                continue
            return text
        except UnicodeDecodeError:
            continue
    raise IngestError("The file's text encoding could not be detected.")


def read_csv(data: bytes, filename: str) -> list[ImportedTable]:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = "\t" if filename.lower().endswith(".tsv") else ","
    grid = [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    table = _from_grid(grid, filename, None, hashlib.sha256(data).hexdigest())
    return [table] if table else []


def read_excel(data: bytes, filename: str) -> list[ImportedTable]:
    sha = hashlib.sha256(data).hexdigest()
    tables = []
    if filename.lower().endswith(".xls"):
        try:
            import xlrd
        except ImportError as exc:
            raise UnsupportedFormatError(f"{filename}: reading .xls needs the 'xlrd' package; "
                                         "or save the sheet as .xlsx/.csv.") from exc
        book = xlrd.open_workbook(file_contents=data)
        for sheet in book.sheets():
            grid = [[clean_cell(sheet.cell_value(r, c)) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
            table = _from_grid(grid, filename, sheet.name, sha)
            if table:
                tables.append(table)
        return tables
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise IngestError(f"{filename} could not be opened as an Excel workbook ({exc.__class__.__name__}).") from exc
    for ws in book.worksheets:
        grid = [[clean_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
        table = _from_grid(grid, filename, ws.title, sha)
        if table:
            tables.append(table)
    book.close()
    return tables


def read_table_file(path: Path | None = None, *, data: bytes | None = None,
                    filename: str | None = None) -> list[ImportedTable]:
    """Read every non-empty sheet of a CSV/Excel file."""
    if data is None:
        data = Path(path).read_bytes()
    filename = filename or Path(path).name
    ext = Path(filename).suffix.lower()
    if ext not in TABLE_EXT:
        raise UnsupportedFormatError(f"{filename}: upload investor lists as CSV or Excel (.xlsx/.xls).")
    if not data.strip():
        raise IngestError(f"{filename} is empty.")
    tables = read_excel(data, filename) if ext in {".xlsx", ".xlsm", ".xls"} else read_csv(data, filename)
    if not tables:
        raise IngestError(f"{filename} has no data rows.")
    return tables
