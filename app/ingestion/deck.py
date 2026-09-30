"""Pitch deck readers: every slide/page, including tables, charts, notes, footnotes and OCR text.

* PDF — PyMuPDF text and tables; pages with little text or large images are rendered and OCR'd.
* PPTX — python-pptx text frames, grouped shapes, tables, chart data and speaker notes;
  pictures (and slides without text) are OCR'd.
* PPT — converted to PPTX with LibreOffice when installed; otherwise a clear instruction.
"""

from __future__ import annotations

import contextlib
import hashlib
import io
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from app.errors import IngestError, UnsupportedFormatError
from app.ingestion import ocr
from app.utils.logging import get_logger
from app.utils.text import normalize_text

LOGGER = get_logger(__name__)
SOFFICE_TIMEOUT_S = 180
PPT_INSTRUCTION = (
    "{name} is a legacy PowerPoint 97-2003 (.ppt) file. Converting it needs LibreOffice, which was not "
    "found. Either install LibreOffice (https://www.libreoffice.org/download — make sure `soffice` is on "
    "PATH or in C:\\Program Files\\LibreOffice\\program), or open the deck in PowerPoint and use "
    "File › Save As › PowerPoint Presentation (.pptx) or PDF, then upload that file."
)


@dataclass
class DeckPage:
    number: int
    text: str = ""
    tables: list[str] = field(default_factory=list)
    notes: str = ""
    ocr_text: str = ""

    @property
    def content(self) -> str:
        parts = [self.text, "\n".join(self.tables), self.notes, self.ocr_text]
        return "\n".join(p for p in parts if p)

    def segments(self) -> list[tuple[str, str]]:
        """(sub-locator, text) pairs, so every fact can cite where on the slide it came from."""
        out = []
        if self.text:
            out.append(("", self.text))
        if self.tables:
            out.append(("table", "\n".join(self.tables)))
        if self.notes:
            out.append(("speaker notes", self.notes))
        if self.ocr_text:
            out.append(("OCR", self.ocr_text))
        return out


@dataclass
class DeckDocument:
    filename: str
    source_format: str
    unit: str                                  # "slide" or "page"
    pages: list[DeckPage]
    sha256: str = ""
    metadata_date: datetime | None = None
    warnings: list[str] = field(default_factory=list)
    ocr_pages: list[int] = field(default_factory=list)
    ocr_skipped: list[int] = field(default_factory=list)

    @property
    def page_count(self) -> int:
        return len(self.pages)

    def locator(self, number: int, segment: str = "") -> str:
        return f"{self.unit} {number}" + (f" ({segment})" if segment else "")

    @property
    def full_text(self) -> str:
        return "\n\n".join(f"[{self.unit.title()} {p.number}]\n{p.content}" for p in self.pages)


class OcrBudget:
    def __init__(self, cfg: dict[str, Any]) -> None:
        self.enabled = bool(cfg.get("enabled", True))
        self.remaining = int(cfg.get("max_pages", 40))
        self.low_text = int(cfg.get("low_text_chars", 40))
        self.image_ratio = float(cfg.get("image_area_ratio", 0.25))
        self.dpi = int(cfg.get("render_dpi", 150))

    def take(self) -> bool:
        if not self.enabled or self.remaining <= 0:
            return False
        self.remaining -= 1
        return True


# ------------------------------------------------------------------------------ PDF


def _pdf_tables(page: Any) -> list[str]:
    rows: list[str] = []
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            finder = page.find_tables()
        for table in finder.tables:
            for row in table.extract():
                cells = [" ".join(str(c).split()) for c in row if c not in (None, "")]
                if len(cells) >= 2:
                    rows.append(" | ".join(cells))
    except Exception:  # noqa: BLE001 - table detection is best effort
        pass
    return rows


def _pdf_layout_rows(page: Any) -> list[str]:
    """Rows of text that sit on one baseline but in separate text lines (tables drawn as positioned text)."""
    segments: dict[tuple[int, int], list] = {}
    try:
        for x0, y0, x1, y1, word, block, line, _n in page.get_text("words"):
            seg = segments.setdefault((block, line), [x0, (y0 + y1) / 2, []])
            seg[0] = min(seg[0], x0)
            seg[2].append(word)
    except Exception:  # noqa: BLE001
        return []
    rows: list[list] = []
    for block_line, (x0, yc, words) in sorted(segments.items(), key=lambda kv: kv[1][1]):
        if rows and abs(rows[-1][0] - yc) <= 3:
            rows[-1][1].append((x0, block_line, " ".join(words)))
        else:
            rows.append([yc, [(x0, block_line, " ".join(words))]])
    out = []
    for _yc, parts in rows:
        if len({b for _x, b, _t in parts}) >= 2:
            out.append(" | ".join(t for _x, _b, t in sorted(parts)))
    return out


def _pdf_image_ratio(page: Any) -> float:
    try:
        area = abs(page.rect)
        covered = sum(abs(info["bbox"] & page.rect) for info in page.get_image_info())
        return min(covered / area, 1.0) if area else 0.0
    except Exception:  # noqa: BLE001
        return 0.0


def _pdf_date(raw: str | None) -> datetime | None:
    if not raw:
        return None
    text = raw[2:] if raw.startswith("D:") else raw
    try:
        return datetime.strptime(text[:8], "%Y%m%d")
    except ValueError:
        return None


def read_pdf(path: Path, display_name: str, ocr_cfg: dict[str, Any]) -> DeckDocument:
    import pymupdf

    try:
        document = pymupdf.open(path)
    except Exception as exc:  # noqa: BLE001
        raise IngestError(f"{display_name} could not be opened as a PDF ({exc.__class__.__name__}).") from exc
    budget = OcrBudget(ocr_cfg)
    ocr_ready = budget.enabled and ocr.available()
    pages: list[DeckPage] = []
    doc = DeckDocument(filename=display_name, source_format="pdf", unit="page", pages=pages)
    with document:
        if document.needs_pass:
            raise IngestError(f"{display_name} is password-protected. Remove the password and upload again.")
        if document.page_count == 0:
            raise IngestError(f"{display_name} contains no pages.")
        doc.metadata_date = _pdf_date((document.metadata or {}).get("creationDate"))
        for index in range(document.page_count):
            page = document[index]
            number = index + 1
            try:
                text = normalize_text(page.get_text("text"))
            except Exception:  # noqa: BLE001
                text = ""
                doc.warnings.append(f"Page {number} text could not be read.")
            tables = [row for row in _pdf_tables(page) if row not in text]
            tables += [row for row in _pdf_layout_rows(page) if row not in text and row not in tables]
            deck_page = DeckPage(number=number, text=text, tables=tables)
            needs_ocr = len(text) < budget.low_text or _pdf_image_ratio(page) >= budget.image_ratio
            if needs_ocr:
                if ocr_ready and budget.take():
                    pix = page.get_pixmap(dpi=budget.dpi)
                    deck_page.ocr_text = normalize_text(ocr.ocr_image(pix.tobytes("png")))
                    doc.ocr_pages.append(number)
                else:
                    doc.ocr_skipped.append(number)
            pages.append(deck_page)
    return doc


# ------------------------------------------------------------------------------ PPTX


def _chart_lines(chart: Any) -> list[str]:
    lines: list[str] = []
    try:
        if chart.has_title and chart.chart_title.has_text_frame and chart.chart_title.text_frame.text.strip():
            lines.append(f"Chart: {chart.chart_title.text_frame.text.strip()}")
        for plot in chart.plots:
            categories = [str(c) for c in plot.categories]
            for series in plot.series:
                values = list(series.values)
                pairs = (", ".join(f"{c}: {v}" for c, v in zip(categories, values, strict=False))
                         if len(categories) == len(values) else ", ".join(map(str, values)))
                lines.append(f"{series.name or 'Series'} | {pairs}")
    except Exception:  # noqa: BLE001 - chart XML varies widely
        pass
    return lines


def _walk_shapes(shapes: Any, slide_area: float, out: dict[str, list]) -> None:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        try:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                _walk_shapes(shape.shapes, slide_area, out)
                continue
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [" ".join(cell.text.split()) for cell in row.cells]
                    if any(cells):
                        out["tables"].append(" | ".join(c for c in cells if c))
                continue
            if getattr(shape, "has_chart", False) and shape.has_chart:
                out["tables"].extend(_chart_lines(shape.chart))
                continue
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame and shape.text_frame.text.strip():
                out["text"].append(shape.text_frame.text)
            image = getattr(shape, "image", None) if shape.shape_type == MSO_SHAPE_TYPE.PICTURE else None
            if image is not None:
                area = float(shape.width or 0) * float(shape.height or 0)
                out["pictures"].append((area / slide_area if slide_area else 0.0, image.blob))
        except Exception:  # noqa: BLE001 - one odd shape must not sink the slide
            continue


def read_pptx(path: Path, display_name: str, ocr_cfg: dict[str, Any]) -> DeckDocument:
    from pptx import Presentation

    try:
        presentation = Presentation(str(path))
    except Exception as exc:  # noqa: BLE001
        raise IngestError(f"{display_name} could not be opened as a PowerPoint file "
                          f"({exc.__class__.__name__}); it may be corrupt or password-protected.") from exc
    budget = OcrBudget(ocr_cfg)
    ocr_ready = budget.enabled and ocr.available()
    slide_area = float(presentation.slide_width or 0) * float(presentation.slide_height or 0)
    pages: list[DeckPage] = []
    doc = DeckDocument(filename=display_name, source_format="pptx", unit="slide", pages=pages)
    try:
        modified = presentation.core_properties.modified or presentation.core_properties.created
        doc.metadata_date = modified if isinstance(modified, datetime) else None
    except Exception:  # noqa: BLE001
        pass
    for number, slide in enumerate(presentation.slides, start=1):
        found: dict[str, list] = {"text": [], "tables": [], "pictures": []}
        _walk_shapes(slide.shapes, slide_area, found)
        text = normalize_text("\n".join(found["text"]))
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = normalize_text(slide.notes_slide.notes_text_frame.text)
        page = DeckPage(number=number, text=text, tables=found["tables"], notes=notes)
        low_text = len(text) + sum(len(t) for t in found["tables"]) < budget.low_text
        targets = [blob for ratio, blob in found["pictures"] if low_text or ratio >= budget.image_ratio]
        if targets:
            if ocr_ready and budget.take():
                page.ocr_text = normalize_text("\n".join(filter(None, (ocr.ocr_image(b) for b in targets))))
                doc.ocr_pages.append(number)
            else:
                doc.ocr_skipped.append(number)
        pages.append(page)
    if not pages:
        raise IngestError(f"{display_name} contains no slides.")
    return doc


# ------------------------------------------------------------------------------ PPT


def find_soffice() -> str | None:
    for candidate in ("soffice", "soffice.exe", "libreoffice"):
        found = shutil.which(candidate)
        if found:
            return found
    for path in (Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
                 Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
                 Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
                 Path("/usr/bin/soffice"), Path("/usr/lib/libreoffice/program/soffice")):
        if path.is_file():
            return str(path)
    return None


def read_ppt(path: Path, display_name: str, ocr_cfg: dict[str, Any]) -> DeckDocument:
    soffice = find_soffice()
    if soffice is None:
        raise UnsupportedFormatError(PPT_INSTRUCTION.format(name=display_name))
    with tempfile.TemporaryDirectory(prefix="im-ppt-") as tmp:
        try:
            result = subprocess.run(  # noqa: S603 - fixed argv, resolved binary
                [soffice, "--headless", "--convert-to", "pptx", "--outdir", tmp, str(path)],
                capture_output=True, timeout=SOFFICE_TIMEOUT_S, check=False,
            )
        except (subprocess.TimeoutExpired, OSError) as exc:
            raise IngestError(f"Converting {display_name} with LibreOffice failed: {exc}") from exc
        produced = sorted(Path(tmp).glob("*.pptx"))
        if result.returncode != 0 or not produced:
            raise IngestError(f"LibreOffice could not convert {display_name}. Save it as .pptx or PDF and retry.")
        doc = read_pptx(produced[0], display_name, ocr_cfg)
    doc.source_format = "ppt"
    doc.warnings.append("Converted from legacy .ppt with LibreOffice; check extracted values against the deck.")
    return doc


# ------------------------------------------------------------------------------ dispatch

_MAGIC = {".pdf": b"%PDF", ".pptx": b"PK\x03\x04", ".ppt": b"\xd0\xcf\x11\xe0"}


def read_deck(path: Path, display_name: str | None = None, ocr_cfg: dict[str, Any] | None = None) -> DeckDocument:
    path = Path(path)
    display_name = display_name or path.name
    ocr_cfg = ocr_cfg or {}
    ext = Path(display_name).suffix.lower() or path.suffix.lower()
    if ext not in _MAGIC:
        raise UnsupportedFormatError(f"{display_name}: upload the pitch deck as PDF, PPTX or PPT.")
    data = path.read_bytes()
    if _MAGIC[ext] not in data[:1024]:
        raise IngestError(f"{display_name} does not look like a valid {ext} file.")
    reader = {".pdf": read_pdf, ".pptx": read_pptx, ".ppt": read_ppt}[ext]
    doc = reader(path, display_name, ocr_cfg)
    doc.sha256 = hashlib.sha256(data).hexdigest()
    if doc.ocr_skipped:
        reason = "OCR is disabled" if not ocr_cfg.get("enabled", True) else (
            "the OCR engine (rapidocr-onnxruntime) is not installed" if not ocr.available()
            else "the OCR page budget was reached")
        doc.warnings.append(f"Image-heavy {doc.unit}s {', '.join(map(str, doc.ocr_skipped[:12]))} were not OCR'd "
                            f"because {reason}; facts shown only in images may be missing.")
    LOGGER.info("Read deck: %d %ss, OCR on %d", doc.page_count, doc.unit, len(doc.ocr_pages))
    return doc
