"""Excel workbook: five sheets with filters, frozen headers, readable widths and formula-safe text.

Text cells are written as literal strings (data type 's') with Excel's quote-prefix flag when they
start with =, +, -, @, tab or CR, so imported values can never execute as formulas — and the value
itself is unchanged (a phone number like "+1 512…" stays exactly as entered).
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models import RunResult
from app.reports.frames import Sheet, build_sheets

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
BLOCK_FONT = Font(bold=True, color="1F3864", size=12)
_RISKY = ("=", "+", "-", "@", "\t", "\r")
MAX_WIDTH, MIN_WIDTH, WRAP_WIDTH = 60, 9, 50


def write_cell(ws: Worksheet, row: int, col: int, value: Any):
    cell = ws.cell(row=row, column=col)
    if isinstance(value, str):
        cell.value = value
        cell.data_type = "s"
        if value.startswith(_RISKY):
            cell.quotePrefix = True
    else:
        cell.value = value
    return cell


def _header(ws: Worksheet, row: int, columns: list[str]) -> None:
    for col, name in enumerate(columns, start=1):
        cell = write_cell(ws, row, col, name)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _write_sheet(ws: Worksheet, sheet: Sheet) -> None:
    _header(ws, 1, sheet.columns)
    for r, row in enumerate(sheet.rows, start=2):
        for c, value in enumerate(row, start=1):
            cell = write_cell(ws, r, c, value)
            name = sheet.columns[c - 1]
            if name in sheet.wrap:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                cell.alignment = Alignment(vertical="top")
            if name in sheet.percent and isinstance(value, (int, float)):
                cell.number_format = "0%"
            if name in sheet.decimal and isinstance(value, (int, float)):
                cell.number_format = "0.0"
    last_row = max(1, len(sheet.rows) + 1)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(sheet.columns))}{last_row}"
    widths = []
    for c, name in enumerate(sheet.columns):
        longest = max([len(str(name))] + [len(str(row[c])) for row in sheet.rows[:500] if row[c] is not None])
        cap = WRAP_WIDTH if name in sheet.wrap else MAX_WIDTH
        widths.append(max(MIN_WIDTH, min(cap, longest + 2)))
    row = last_row + 2
    for title, cols, rows in sheet.extra_blocks:
        write_cell(ws, row, 1, title).font = BLOCK_FONT
        _header(ws, row + 1, cols)
        for i, values in enumerate(rows, start=row + 2):
            for c, value in enumerate(values, start=1):
                write_cell(ws, i, c, value).alignment = Alignment(wrap_text=True, vertical="top")
        for c in range(min(len(cols), len(widths))):
            longest = max([len(str(v[c])) for v in rows if c < len(v)] + [len(cols[c])])
            widths[c] = max(widths[c], min(WRAP_WIDTH, longest + 2))
        row += len(rows) + 3
    for c, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = width


def render_xlsx(result: RunResult) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in build_sheets(result):
        _write_sheet(wb.create_sheet(sheet.name), sheet)
    wb.properties.title = f"Investor Match — {result.context.company_name or 'company'}"
    wb.properties.creator = "TEN Capital Network"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
